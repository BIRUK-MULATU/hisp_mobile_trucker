#!/usr/bin/env python3
"""Builds the RDHIS2 Mobile — Technical Documentation .docx (and, if libreoffice
is on PATH, a PDF next to it).

    pip install --break-system-packages python-docx
    python3 scripts/build_technical_doc.py

Output: technical_documentation/RDHIS2_Mobile_Technical_Documentation.{docx,pdf}

This is the stakeholder-facing companion to docs/ — keep it in step with the
docs/ chapters after significant changes.
"""
import datetime as dt
import os
import shutil
import subprocess
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BUILD_DATE = dt.date.today()
LONG_DATE = BUILD_DATE.strftime('%d %B %Y')
ISO_DATE = BUILD_DATE.isoformat()

NAVY = RGBColor(0x1E, 0x28, 0x44)
ACCENT = RGBColor(0x28, 0x5A, 0x9E)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ---- base styles -----------------------------------------------------------
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

for lvl, sz, col in [(1, 17, NAVY), (2, 13.5, NAVY), (3, 11.5, ACCENT)]:
    st = styles[f'Heading {lvl}']
    st.font.name = 'Calibri'
    st.font.size = Pt(sz)
    st.font.color.rgb = col
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

def _set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def para(text="", *, bold=False, italic=False, size=None, color=None,
         align=None, style=None, space_after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p

def rich(parts, *, style=None, align=None, space_after=None):
    """parts: list of (text, {kwargs}) tuples."""
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    for text, kw in parts:
        r = p.add_run(text)
        r.bold = kw.get('bold', False)
        r.italic = kw.get('italic', False)
        if kw.get('mono'):
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
        if kw.get('size'):
            r.font.size = Pt(kw['size'])
        if kw.get('color'):
            r.font.color.rgb = kw['color']
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p

def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def h3(t): doc.add_heading(t, level=3)

def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p

def bullet_rich(parts, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    for text, kw in parts:
        r = p.add_run(text)
        r.bold = kw.get('bold', False)
        r.italic = kw.get('italic', False)
        if kw.get('mono'):
            r.font.name = 'Consolas'; r.font.size = Pt(9.5)
    p.paragraph_format.space_after = Pt(3)
    return p

def numbered(text):
    p = doc.add_paragraph(style='List Number')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p

def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_cell_bg(hdr[i], '1E2844')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            para0 = cells[i].paragraphs[0]
            run = para0.add_run(str(val))
            run.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:color'), 'BBBBBB')
    bdr.append(bottom)
    pPr.append(bdr)

# ===========================================================================
# TITLE PAGE
# ===========================================================================
for _ in range(4):
    doc.add_paragraph()
para("RDHIS2 Mobile", bold=True, size=34, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Technical Documentation", size=20, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
para("Offline-first Flutter application for aggregate health-data capture into DHIS2",
     italic=True, size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=48)
para("HISP Ethiopia  ·  in collaboration with the Federal Ministry of Health",
     size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para(f"Document version 1.0  ·  {LONG_DATE}",
     size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Audience: engineering, technical reviewers, HISP / Ministry stakeholders",
     size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ===========================================================================
# DOCUMENT CONTROL
# ===========================================================================
h1("Document Control")
table(["Field", "Value"],
      [["Title", "RDHIS2 Mobile — Technical Documentation"],
       ["Version", "1.0"],
       ["Date", LONG_DATE],
       ["Status", "Baseline — reflects the codebase at the current development head"],
       ["Owner", "HISP Ethiopia — Mobile team"],
       ["Application", "RDHIS2 Mobile App (package: hisp_mobile_trucker), app version 1.0.0+1"],
       ["Source of truth", "The repository. Every claim here is traceable to a file "
        "path; where behaviour and this document disagree, the code is authoritative."],
       ["Companion docs", "docs/ (16+ developer-onboarding chapters); "
        "OFFLINE_INTEGRATION.md (historical design note); "
        "ministry_presentation/ (stakeholder deck)"]],
      widths=[1.6, 4.6])

para("Scope. ", bold=True)
para("This document describes what the application is, how it is built, how the "
     "offline-first model works end to end, the data-quality and synchronization "
     "machinery, the security posture, the database schema, and the build / release / "
     "test process. It is written to be read by a newcomer to the project and by a "
     "technical reviewer who will not contribute code but needs an accurate picture "
     "of the system.")

para("Change history.", bold=True)
table(["Version", "Date", "Notes"],
      [["1.0", "2026-09-10", "First consolidated technical documentation. Supersedes "
        "the point-in-time notes in ARCHITECTURE_ANALYSIS.md and folds in features "
        "added since the docs/ set was written (data-quality suite, outlier detection, "
        "audit trail, schema migrations, local dashboards, deadline reminders, "
        "onboarding, background-sync hardening)."]],
      widths=[0.8, 1.1, 4.3])
doc.add_page_break()

# ===========================================================================
# 1. EXECUTIVE SUMMARY
# ===========================================================================
h1("1  Executive Summary")

para("RDHIS2 Mobile is an Android / iOS / web application, built in Flutter, that lets "
     "health workers record aggregate service-delivery data and submit it to DHIS2 — "
     "the national Health Management Information System. It is designed for the "
     "conditions the data is actually collected in: facilities with no connectivity, "
     "intermittent connectivity, or a connection too weak to depend on.")

h3("The one-sentence mental model")
rich([("Every screen reads from and writes to a ", {}),
      ("local, per-user SQLite database first.", {'bold': True}),
      (" The network is not a dependency for daily work — it is how that local "
       "database is seeded (metadata sync) and how locally-queued entries eventually "
       "reach the DHIS2 server (data sync). ", {}),
      ("The local database is the truth for the UI; the server is a sync target, "
       "not a live dependency.", {'bold': True})])

h3("What it does")
bullet("Online and offline login. The first login is verified against the DHIS2 "
       "server; afterwards the same credentials work offline indefinitely, checked "
       "locally against a salted hash that never stores the password.")
bullet("Offline-first capture. Organisation-unit tree → dataset → section → period → "
       "entry form. Every keystroke is saved locally. Nothing is lost when the "
       "network drops mid-form.")
bullet("Draft → Complete workflow. Entries stay as device-only drafts until the user "
       "marks the form Complete; only then are they queued for the server. Completed "
       "forms can be reopened for correction.")
bullet("Automatic and manual synchronization, with per-value server verdicts. Values "
       "the server rejects are shown in red in the form with the server's own reason; "
       "editing them re-queues the fix.")
bullet("Data quality at the point of entry: client-side value-type validation, "
       "offline evaluation of the server's validation rules, mandatory-field "
       "enforcement, grey (disabled) fields, controller-element gating, and a "
       "statistical outlier check against the cell's own recent history.")
bullet("A per-cell audit trail combining this device's local edit history with "
       "DHIS2's own server-side audit log.")
bullet("Native dashboards. The Visualization area renders DHIS2 server dashboards "
       "with Flutter charts, lets the user build and save local dashboards on the "
       "device, and caches results so a chart is still viewable offline.")
bullet("Ethiopian-calendar period handling throughout, including the DHIS2 "
       "Ethiopian-fiscal-year period types.")
bullet("On-device deadline reminders for reports the user still owes, and a "
       "first-run onboarding tour.")

h3("Project size (current head)")
table(["Metric", "Value"],
      [["Production Dart", "~29,000 lines across ~160 files (excludes generated code)"],
       ["Generated code", "~16,900 lines (Drift database, committed to the repo)"],
       ["Automated tests", "24 test files, ~5,200 lines — weighted toward the "
        "offline-critical core (sync, conflict resolution, validation, calendar)"],
       ["Platforms", "Android (primary), iOS, Web (development / demo)"],
       ["Toolchain", "Flutter (CI pins 3.41.4), Dart SDK >= 3.0.0 < 4.0.0"]],
      widths=[1.7, 4.5])

h3("Security posture in one paragraph")
para("Passwords are never stored in any recoverable form, online or offline — offline "
     "login is checked against a salted SHA-256 verifier bound to the server URL, "
     "compared in constant time. Session and credential material live in OS-backed "
     "secure storage (Android Keystore-backed EncryptedSharedPreferences / iOS "
     "Keychain), never in the SQLite database. Transport is HTTPS-only in release "
     "builds. The known hardening gaps — local database not encrypted at rest (it "
     "holds cached metadata and pending field data, never credentials), no "
     "certificate pinning, no R8/ProGuard — are defense-in-depth items scoped to the "
     "“lost or stolen device” threat model, each with a concrete fix on the "
     "roadmap. Full detail in section 9.")
doc.add_page_break()

# ===========================================================================
# 2. SYSTEM OVERVIEW
# ===========================================================================
h1("2  System Overview")

h2("2.1  Purpose and context")
para("DHIS2 is the platform of record for routine health data in Ethiopia. Its web "
     "and official Android clients assume a usable connection at the moment of data "
     "entry. In much of the country that assumption does not hold: a health worker "
     "may capture a full day or month of data with no connectivity and only reach a "
     "network hours or days later. RDHIS2 Mobile exists to make that workflow safe — "
     "capture now, sync later, lose nothing, and never silently overwrite work.")

h2("2.2  Users and roles")
bullet("Facility data-entry users — the primary audience. Assigned a small "
       "organisation-unit scope (their facility, sometimes a woreda and its direct "
       "children) and a set of datasets. They capture, complete, and sync reports.")
bullet("Supervisory / woreda users — a wider org-unit scope; the same capture flow "
       "plus dashboards spanning their subtree.")
para("There is no in-app role management: the app inherits whatever the DHIS2 server "
     "grants the authenticated user (org-unit capture scope, dataset access, "
     "metadata visibility). The app never widens that scope.")

h2("2.3  DHIS2 integration surface")
para("The app speaks the standard DHIS2 Web API over HTTP Basic authentication. The "
     "endpoints it uses, and nothing else:")
table(["Purpose", "Endpoint(s)", "Direction"],
      [["Credential check on login", "GET /api/me.json", "read"],
       ["Metadata sync", "GET /api/{resource}.json (categories, category "
        "combos/options, option sets, attributes, data elements, indicators, data "
        "element groups, organisation units, data sets, sections, validation rules)",
        "read"],
       ["Form open — pull current values", "GET /api/dataValueSets.json", "read"],
       ["Outlier history snapshot", "GET /api/dataValueSets.json (date-windowed)", "read"],
       ["Data push", "POST /api/dataValueSets.json (importStrategy "
        "CREATE_AND_UPDATE, atomicMode NONE)", "write"],
       ["Completion push", "POST / DELETE /api/completeDataSetRegistrations", "write"],
       ["Server audit trail", "GET /api/audits/dataValue.json", "read"],
       ["Connectivity probe", "GET /api/system/ping (or {baseUrl}/system/ping)", "read"],
       ["Dashboards & analytics", "GET /api/dashboards, /api/visualizations/{id}, "
        "/api/analytics, /api/indicatorGroups, /api/dataElementGroups", "read"]],
      widths=[1.6, 3.4, 0.7])
para("The app never creates or edits a server-side Visualization, dashboard, "
     "metadata object, or user. Its only writes are data values and completion "
     "registrations.", italic=True)

h2("2.4  Server configuration")
rich([("The compiled-in default server is ", {}),
      ("https://hmis-staging.moh.gov.et/api", {'mono': True}),
      (" (", {}), ("lib/core/constants/api_constants.dart", {'mono': True}),
      ("). Users and developers can repoint the app at any DHIS2 instance at "
       "runtime — no rebuild — from the server icon on the login screen or "
       "Settings → DHIS2 server URL. The override is normalised to "
       "https:// + /api and persisted in secure storage.", {})])
doc.add_page_break()

# ===========================================================================
# 3. ARCHITECTURE
# ===========================================================================
h1("3  Architecture")

h2("3.1  Style: feature-first Clean Architecture")
para("Each feature under lib/features/ owns its own data / domain / presentation "
     "split. Cross-cutting infrastructure lives in lib/core/; shared UI in "
     "lib/shared/. The three-layer split is applied where there is a real "
     "repository / use-case boundary to draw, not as ceremony: home and settings are "
     "presentation-only because they compose other features and own no persistence.")

h3("The four layers, and why each exists")
table(["Layer", "Contents", "Rationale"],
      [["Presentation", "Pages, widgets, Bloc classes. Owns UI state, translates "
        "user intent into domain calls. Never touches Dio or Drift directly.",
        "A screen can be redesigned or rebuilt without touching business rules."],
       ["Domain", "Entities, use cases, repository interfaces. Pure Dart — no "
        "Flutter, Dio, or Drift imports anywhere in a domain folder.",
        "Business rules are testable in total isolation and would survive a rewrite "
        "of the UI, the network stack, or the database."],
       ["Data", "Repository implementations, remote/local data sources, models "
        "(DTOs with fromJson/toJson that extend the domain entities).",
        "One place per feature holds its persistence strategy. Auth and data-entry "
        "each migrated from remote-only to offline-first here without a single Bloc "
        "or page changing."],
       ["Core", "HTTP client, database, secure storage, sync engine, session/auth, "
        "shared exception types, the Ethiopian calendar, data-quality services.",
        "One source of truth for “is the user logged in”, “are we "
        "online”, “what period is it” — no feature reinvents them."]],
      widths=[1.0, 2.7, 2.5])

h3("Dependency direction")
code_block("Presentation ──▶ Domain ◀── Data\n"
           "                   ▲\n"
           "                   └──  Core  (network, database, storage, sync, auth, errors)")
para("Presentation and Data both depend on Domain; Domain depends on nothing "
     "feature-specific. A Bloc holds a repository interface, never the concrete "
     "implementation — which is what makes it possible to inject a fake in a test.")

h2("3.2  Module map")
code_block(
"lib/\n"
"├── main.dart                 App entry: wires singletons, starts sync, loads flags\n"
"├── core/\n"
"│   ├── auth/                 AppSession, SessionService (the login decision tree),\n"
"│   │                         CredentialStore (offline verifier)\n"
"│   ├── constants/            API + app constants\n"
"│   ├── data/                 DataValueStore/Sync/Push, Completeness, AuditLogStore,\n"
"│   │                         ServerAuditService, ValidationService,\n"
"│   │                         OutlierDetectionService, ControllerElementService,\n"
"│   │                         element/indicator label services,\n"
"│   │                         Ethiopian calendar + period service, PeriodAccess\n"
"│   ├── database/             Drift schema + per-platform connection\n"
"│   ├── errors/               AppException (data layer) + Failure (domain layer)\n"
"│   ├── metadata/             MetadataResource base + one file per DHIS2 type\n"
"│   ├── network/              ApiClient (Dio), interceptors, ConnectivityService\n"
"│   ├── notifications/        ReportReminderService (deadline reminders)\n"
"│   ├── onboarding/           OnboardingService, TourHelper (first-run flags)\n"
"│   ├── router/               go_router config + reactive auth guard\n"
"│   ├── storage/              SecureStorage wrapper\n"
"│   └── sync/                 SyncManager, DriftSyncManager, SyncCoordinator,\n"
"│                             manual sync, SyncForegroundService, BatteryOptimization\n"
"├── debug/                    Dev-only screens (sync debug, in-app DB viewer)\n"
"├── features/\n"
"│   ├── auth/                 Login (online/offline)\n"
"│   ├── capture/              Org-unit tree → dataset → section → period; Report\n"
"│   │                         Period view; Expected-reports “To-do”; new-report flow\n"
"│   ├── data_entry/           The entry form: bloc, table, cells, outlier + validation\n"
"│   │                         UI, disease-registration list, PDF/Excel export\n"
"│   ├── audit_log/            Per-cell history sheet (local + server audit merged)\n"
"│   ├── home/                 App shell: Capture / Visualization toggle, filters, drawer\n"
"│   ├── onboarding/           First-run carousel\n"
"│   ├── settings/             Server URL, profile, logout\n"
"│   └── visualization/        Server dashboards, local dashboards, chart builder\n"
"└── shared/\n"
"    ├── theme/                AppColors, AppTextStyles, AppDimensions, AppTheme\n"
"    └── widgets/              Buttons, fields, loaders, FilterPanel, SegmentedToggle,\n"
"                              ServerUrlDialog, SyncSnackbar, squircle FAB, option grid")

h2("3.3  Technology choices")
table(["Concern", "Choice", "Version"],
      [["UI framework", "Flutter / Dart", "SDK >= 3.0.0 < 4.0.0; CI Flutter 3.41.4"],
       ["State management", "flutter_bloc (Bloc pattern, no Cubit)", "^8.1.6"],
       ["Navigation", "go_router with a reactive login guard", "^14.2.7"],
       ["HTTP", "Dio (background JSON transformer)", "^5.7.0"],
       ["Local database", "Drift over SQLite (sqlite3_flutter_libs native, "
        "sqlite3.wasm on web)", "drift ^2.20.0"],
       ["Secure storage", "flutter_secure_storage", "^9.2.2"],
       ["Offline credential check", "crypto (SHA-256)", "^3.0.5"],
       ["Connectivity", "connectivity_plus + an active server probe", "^7.2.0"],
       ["Charts", "fl_chart (native rendering)", "^1.2.0"],
       ["Reminders", "flutter_local_notifications + timezone", "^17.2.3 / ^0.9.4"],
       ["Onboarding", "showcaseview + shared_preferences", "^4.0.1 / ^2.3.2"],
       ["Export", "pdf + printing + excel + share_plus", "—"],
       ["Dev tooling", "build_runner, drift_dev, flutter_lints, drift_db_viewer", "—"]],
      widths=[1.5, 3.0, 1.7])

h2("3.4  Dependency injection")
para("There is deliberately no DI framework (no get_it, riverpod, injectable, "
     "provider). App-wide singletons use the classic private-constructor + static "
     "instance pattern (ApiClient, SecureStorage, AppSession.instance, "
     "ConnectivityService.instance, DriftSyncManager.instance). Feature wiring is "
     "done by hand at page build time. Repository implementations take their "
     "dependencies (session, network info) as optional named parameters that default "
     "to the real singleton — which is what makes them constructible with fakes in "
     "tests without a container.")
rich([("Trade-off, stated plainly: ", {'bold': True}),
      ("this is fully traceable (cmd-click from a widget to every dependency, no "
       "reflection or codegen) at the cost of the same wiring snippet repeating "
       "across pages. A small service locator is the natural next step if the "
       "feature count keeps growing — it is on the roadmap.", {})])

h2("3.5  Startup sequence (main.dart)")
para("On launch, in order: bind the Flutter engine; load onboarding / app-tour "
     "flags synchronously into memory (the router's redirect is a synchronous "
     "function and must be able to read them); lock portrait orientation; apply a "
     "Settings-stored server URL over the compiled default; start the SyncCoordinator "
     "(auto-sync, no-ops while logged out); start the ConnectivityService probe; "
     "start the chart-draft coordinator (finishes offline-saved chart drafts when "
     "connectivity returns); initialise the report-reminder service (no-op on web, "
     "self-disables if the platform refuses); then run the app. There is no splash "
     "screen and no explicit auth check here — the router's redirect combined with "
     "AppSession as a refreshListenable handles that reactively.")
doc.add_page_break()

# ===========================================================================
# 4. OFFLINE-FIRST DESIGN
# ===========================================================================
h1("4  Offline-First Design")

h2("4.1  The per-user local database")
rich([("AppDatabase.forUser(userKey)", {'mono': True}),
      (" opens a SQLite connection scoped to a sanitised username "
       "(lower-cased, non-alphanumerics collapsed to underscore). Native builds get "
       "one file per user at ", {}),
      ("<app documents>/hisp_<userKey>.sqlite", {'mono': True}),
      ("; web builds get a WASM-backed database of the same name (OPFS, falling back "
       "to IndexedDB). A single conditional export picks the platform "
       "implementation.", {})])
rich([("Why per-user files matter: ", {'bold': True}),
      ("a shared device with several users has fully isolated metadata caches "
       "and fully isolated pending field data. There is no cross-user query surface "
       "at all — one user's drafts cannot leak into another's session because they "
       "are not in the same database. (The one known edge case: two usernames that "
       "differ only in characters the sanitiser strips would collide.)", {})])

h2("4.2  The write-path state machine")
para("Every user-entered value moves through exactly these states, never skipping "
     "one. The states are stored by enum index — new states may only be appended.")
table(["State", "Meaning", "Sync behaviour"],
      [["draft", "Device-only. Written by every ordinary Save.",
        "No sync path — automatic or manual — ever touches a draft row."],
       ["pending", "Sync-eligible. Set when a form is Completed (drafts are "
        "promoted), or by conflict resolution when the local side wins.",
        "Pushed on the next sync attempt; stays pending on transport failure."],
       ["synced", "Confirmed accepted by the server.",
        "Skipped by future pushes. Eligible for retention purge if out of window."],
       ["error", "Rejected by the server; the reason is stored alongside.",
        "Shown red in the form. Editing the cell clears the error and re-queues "
        "it as pending."]],
      widths=[0.8, 2.7, 2.7])
code_block(
"user types in a cell\n"
"        │  Save\n"
"        ▼\n"
"   [draft] ── device only, no network ─────────────────────────────\n"
"        │  Complete  →  promoteDrafts()\n"
"        ▼\n"
"   [pending] ── push (immediate best-effort, or later via SyncCoordinator)\n"
"        │\n"
"        ├─ server accepts ─────▶ [synced]\n"
"        └─ server rejects ─────▶ [error] + reason   ── edit cell ──▶ [pending]")
rich([("The core offline-safety guarantee: ", {'bold': True}),
      ("work in progress (a draft) cannot be pushed half-finished, and cannot be "
       "silently overwritten by a conflicting server value while the user is still "
       "editing it.", {})])

h2("4.3  Conflict resolution — per cell, on form open")
para("Whenever a form is opened online, DataValueSync.syncForm() pulls the server's "
     "current values for that dataset / period / org unit and resolves each cell "
     "before the form renders. This is best-effort: a pull failure just means the "
     "local data stands.")
table(["Situation", "Resolution"],
      [["No local value, or local already synced", "Server value applied as-is, "
        "marked synced."],
       ["Local value equals server value", "Marked synced, not re-transmitted."],
       ["Local value is a draft", "Draft always wins. Unsubmitted device-only work "
        "is never silently replaced; it leaves the device only when the user "
        "completes the form."],
       ["Local is pending/error and differs from server — a real conflict",
        "Device clock flagged as tampered → server wins. No server timestamp → "
        "server wins. Local lastModified is after (server lastUpdated + 2 min "
        "tolerance) → local wins (stays pending, pushed next). Otherwise → server "
        "wins."]],
      widths=[2.6, 3.6])
para("The 2-minute tolerance is a conservative bias toward the server for near-ties "
     "— it absorbs residual clock skew without letting a genuinely stale local edit "
     "win on a technicality. After resolution, everything still pending for the form "
     "is pushed; only once nothing local remains pending anywhere is the app clock "
     "re-anchored to the server's Date header.")

h2("4.4  The tamper-resistant clock")
para("Period-based access control (“can I still edit last month's report?”) "
     "is meaningless if the device clock can be wound back. PeriodAccess defends it:")
bullet("High-water mark. The latest time the app has ever observed, persisted. "
       "effectiveNow() never returns earlier than this mark — moving the clock back "
       "literally cannot make it go backward.")
bullet("Tamper flag. A backward jump beyond a 2-minute tolerance (to tolerate NTP "
       "corrections) sets a persistent flag. While set: entry into past periods is "
       "refused/warned; entry into the current open period continues (the data is "
       "not suspect, only backdating claims are).")
bullet("Healing. Any successful online contact anchors the clock to the server's "
       "Date header and clears the flag — gated so it will not re-anchor underneath "
       "data that is still pending.")
bullet("Every local timestamp (data-value lastModified, completion date) comes from "
       "effectiveNow(), not raw DateTime.now() — so a backdated device can never "
       "produce a timestamp earlier than something the app has already seen, which "
       "is the property the newest-wins rule depends on.")

h2("4.5  Retention")
rich([("purgeOutsideRetention(allowedPeriods)", {'mono': True}),
      (" deletes rows that are ", {}), ("both", {'italic': True}),
      (" synced ", {'italic': True}),
      ("and outside the allowed period set. Rows in pending, draft, or error state "
       "are never auto-deleted, regardless of period. A device is never allowed to "
       "silently destroy field work that has not confirmed as synced.", {})])
doc.add_page_break()

# ===========================================================================
# 5. DATA CAPTURE & QUALITY
# ===========================================================================
h1("5  Data Capture and Quality")

h2("5.1  The capture flow")
code_block(
"Home ── Capture mode\n"
"  → Organisation-unit tree  (lazily loaded one level at a time; ~38k units on the\n"
"    national instance are never loaded whole)\n"
"  → Datasets assigned to the selected org unit  (Routine + Disease Registration,\n"
"    one merged list; each annotated synced / unsynced)\n"
"  → Section picker  (only if the dataset has sections)\n"
"  → Period picker   (gated by PeriodAccess: expiryDays / openFuturePeriods / clock)\n"
"  → [ Disease datasets: a category-combo picker, e.g. Department × Outcome ]\n"
"  → Entry form")
para("The entire path down to the form is pure local database reads — no network "
     "call happens in the capture feature. The network first appears when the form "
     "opens (a best-effort value pull) and again at sync time.")
rich([("The “Create new record” entry point ", {'bold': True}),
      ("(the + FAB) is a shortcut into the same flow for starting a fresh report "
       "without walking the tree.", {})])

h2("5.2  Ethiopian calendar and period IDs")
para("Two files carry all date/period logic. EthiopianCalendar does pure calendar "
     "math via a Julian Day Number (integer day count), a fixed epoch offset, and a "
     "4-year leap cycle; every month 1–12 has 30 days, month 13 (Pagume) has 5 or 6. "
     "It generates DHIS2-ready period IDs for every period type.")
rich([("The one non-obvious part — the “Nov” period types. ", {'bold': True}),
      ("On a DHIS2 server running the Ethiopian calendar, “Nov” in a "
       "period-type name does not mean Gregorian November — it means Ethiopian month "
       "11 (Hamle). These are DHIS2's Ethiopian-fiscal-year period types (fiscal "
       "year = Hamle–Sene) and their ID carries the year the fiscal year ", {}),
      ("ends", {'italic': True}),
      (" in. Verified against staging analytics.", {})])
rich([("EthiopianCalendar.today()", {'mono': True}),
      (" is the single source of truth for “what period is current” "
       "throughout the app, independently verified against the standard (Amete "
       "Mihret) calendar. Note: the staging server's own period generation has been "
       "observed running ~2 months behind the real calendar — that is a server-side "
       "issue, not an app bug (see section 13).", {})])

h2("5.3  Client-side value-type validation")
rich([("validateDataValue()", {'mono': True}),
      (" mirrors DHIS2's own server-side valueType checks (NUMBER, INTEGER, "
       "PERCENTAGE, UNIT_INTERVAL, BOOLEAN, DATE, PHONE_NUMBER, EMAIL, option-set "
       "membership, …) and runs before a value is ever queued, so malformed data "
       "never leaves the device. It is also a second gate inside the data-entry "
       "Bloc's save path: every user-edited cell is re-validated before the save is "
       "allowed through. Unknown or free-text types are accepted as-is — the server "
       "stays the final authority.", {})])

h2("5.4  Validation rules (offline)")
para("ValidationService runs the synced DHIS2 validation rules against one form "
     "instance's local values, fully offline, on exactly the data the user sees. "
     "It contains a small expression evaluator implementing the same subset the "
     "official DHIS2 Android SDK evaluates offline: arithmetic over data-element "
     "operands, numbers, parentheses. Anything beyond that (d2 functions, constants, "
     "indicators, [days], org-unit groups) causes the rule to be skipped rather than "
     "guessed at.")
bullet("Only rules whose period type matches the dataset and that reference at least "
       "one element of the form are considered.")
bullet("Missing-value strategies (NEVER_SKIP / SKIP_IF_ANY_VALUE_MISSING / "
       "SKIP_IF_ALL_VALUES_MISSING) are honoured; the pair operators "
       "(compulsory_pair, exclusive_pair) are evaluated on presence, not arithmetic.")
bullet_rich([("Violations warn, they never block completion", {'bold': True}),
             (" — the same contract as the DHIS2 web and Android clients. Each "
              "violation shows the rule name, a human-readable comparison "
              "(“ANC 1st visit (12) should be ≤ ANC follow-up (10)”), "
              "the author's instruction text, and importance.", {})])
bullet("Runs live as the user types (against Bloc state) and again at completion "
       "(against the saved local values).")

h2("5.5  Mandatory fields")
rich([("Compulsory ", {}),
      ("dataSetElement", {'mono': True}),
      (" fields (DHIS2 ", {}),
      ("dataSetElement.compulsory", {'mono': True}),
      (") and compulsory element+combo operand pairs "
       "(", {}),
      ("dataSet.compulsoryDataElementOperands", {'mono': True}),
      (") are synced and enforced. Unlike validation rules, a missing mandatory "
       "field ", {}),
      ("does block completing", {'bold': True}),
      (" the form — the check runs dataset-wide (not just the current section), "
       "again matching the DHIS2 web / Android contract.", {})])

h2("5.6  Grey fields and controller elements")
bullet_rich([("Grey (disabled) fields", {'bold': True}),
             (" — per-section disabled cells (DHIS2 ", {}),
             ("greyedFields", {'mono': True}),
             (") are synced and rendered non-editable, exactly as the server "
              "configures them.", {})])
bullet_rich([("Controller data elements", {'bold': True}),
             (" — a Boolean data element tagged with the custom attribute "
              "“Controller Data Element Attribute” = true gates the "
              "visibility of the other elements in its data-element group. Example: "
              "“RMNCH – Delivery services provided” (Yes/No) shows or hides "
              "every other element in its “Delivery Services” group; "
              "answering No hides them and their entered data is cleared. Purely a "
              "metadata lookup — no network.", {})])

h2("5.7  Outlier detection")
para("A statistical companion to the rule check: as the user types a numeric value, "
     "it is judged against that same cell's own recent history at that org unit.")
bullet_rich([("History snapshot. ", {'bold': True}),
             ("OutlierDetectionService.fetchHistory pulls ~25 months of the "
              "dataset's values for the org unit (so a monthly dataset has two full "
              "years of the same season to compare against), reduces each cell to "
              "summary statistics, and caches the result. Offline, it falls back to "
              "the last cached snapshot. An empty result means “check "
              "disabled” — never an error.", {})])
bullet_rich([("The server's own /api/outlierDetection endpoint is deliberately "
              "not used", {'bold': True}),
             (" — it only scores values already stored server-side and cannot "
              "judge the value being typed right now.", {})])
bullet_rich([("Algorithms (user-selectable, one global setting). ", {'bold': True}),
             ("Modified Z-score (median / MAD — robust to the outliers it is "
              "looking for; the default), Z-score (mean / standard deviation), and "
              "Min–Max (outside the historical range plus a guard band). Threshold "
              "is configurable; the default is 3.", {})])
bullet_rich([("Behaviour. ", {'bold': True}),
             ("A flagged value raises a warning dialog with the typical value, the "
              "historical range, and the acceptable bounds. The user can keep the "
              "value (acknowledged once, not re-prompted) or correct it. Values with "
              "too little history (n < 3) or no spread are never flagged. Settings "
              "persist in SharedPreferences.", {})])

h2("5.8  Audit trail")
para("The per-cell “history” sheet merges two sources into one timeline:")
table(["Source", "What it is", "Scope"],
      [["AuditLogStore (local)", "Append-only trail of edits made on this device — "
        "old value, new value, who, when. Written by DataValueStore.setValue and "
        "CompletenessStore.setComplete themselves, so every write path is covered.",
        "This device only. Works offline."],
       ["ServerAuditService (server)", "DHIS2's own audit trail via "
        "GET /api/audits/dataValue — authoritative history reflecting edits from "
        "every client (web, other devices). CREATE / UPDATE / DELETE per entry.",
        "Server-wide. Requires data-value auditing enabled server-side; an empty "
        "result does not prove nothing changed."]],
      widths=[1.5, 3.3, 1.4])
para("Both are keyed to the same field shape so the two sources render identically "
     "in the sheet. A no-op is recorded for an edit that changes nothing.")

h2("5.9  Form export")
para("An open form can be exported to PDF (via the platform print/share sheet) or "
     "to an Excel workbook (shared as a file). Useful for a paper backup or for "
     "supervisory review off-device.")
doc.add_page_break()

# ===========================================================================
# 6. SYNCHRONIZATION
# ===========================================================================
h1("6  Synchronization and Connectivity")

h2("6.1  Two connectivity concepts, on purpose")
table(["Mechanism", "Answers", "Used by"],
      [["NetworkInfo (connectivity_plus wrapper)", "“Does the OS report a "
        "network interface?”", "SyncCoordinator (when to attempt a push); "
        "repositories (whether to try a best-effort server call before falling back "
        "to local)."],
       ["ConnectivityService (active probe)", "“Can we actually reach the "
        "server?” — probes GET {baseUrl}/system/ping every 30 s and on every OS "
        "connectivity event. Any HTTP response (including 401) counts as online; "
        "only a transport failure counts as offline.",
        "The user-facing online/offline indicator. Uses a bare Dio client with no "
        "auth interceptor so a 401 here never ends the session."]],
      widths=[1.9, 2.6, 1.7])

h2("6.2  Automatic sync — three doors, one push")
para("SyncCoordinator.start() (called once in main) listens for three independent "
     "triggers, any of which calls the same idempotent push:")
numbered("Offline → online transition (a connectivity change where the previous "
         "state was disconnected).")
numbered("Login while already online (a previous offline session may have queued "
         "work with no connectivity event to catch it).")
numbered("A 5-minute heartbeat (catches pushes stuck behind a connection that looks "
         "alive — captive portal, flaky link — that would otherwise wait forever).")
para("pushPending() is safe to call repeatedly (synced rows are skipped) and no-ops "
     "while logged out. Failures are swallowed with a debug log — a background sync "
     "must never crash the app; failed rows stay queued for the next attempt.")

h2("6.3  Manual sync")
para("The sync button on Home / Capture runs a precise, user-facing flow: count "
     "pending work first; force-refresh connectivity (the cached flag can be 30 s "
     "stale); if offline, return “will sync automatically”; if nothing "
     "pending, return “already synced” (mentioning any drafts still on "
     "device); otherwise push, then best-effort refresh metadata, and report one of "
     "offline / nothingToSync / uploaded / partial / failed, each with a specific "
     "message.")

h2("6.4  The push and its server verdict")
bullet("Data values go out as POST /api/dataValueSets.json with importStrategy "
       "CREATE_AND_UPDATE and atomicMode NONE, capped at 500 values per request and "
       "chunked recursively — a mid-transfer drop on a flaky link only loses the "
       "unsent slice, not the whole backlog.")
bullet_rich([("DHIS2 2.38+ answers an import with rejected values with HTTP 409, but "
              "the body is still a full ImportSummary. The push treats 200 and 409 "
              "identically once it has that body: each conflict's index maps back to "
              "the exact value in the payload — an ", {}),
             ("exact per-value verdict, not a guess", {'bold': True}),
             (". Values not listed as rejected are marked synced; listed ones get "
              "error with the server's message. A true transport failure (no "
              "parseable summary) marks nothing — every value stays pending.", {})])
bullet_rich([("Cleared cells are pushed as ", {}),
             ("deleted: true", {'mono': True}),
             (", not ", {}),
             ("value: \"\"", {'mono': True}),
             (" (which DHIS2 rejects with E7610). A pre-push screening pass also "
              "settles type-invalid values locally before the POST.", {})])
bullet("Completion registrations push one call each: completed → POST, "
       "un-completed → DELETE. A 409 here is treated as a permanent rejection "
       "(retrying an already-rejected registration forever cannot succeed).")
bullet_rich([("Stale dataset↔org-unit assignment (E7629) is recovered "
              "automatically: ", {}),
             ("the picker refreshes and replaces the assignment on every online "
              "load, and a requeue pass un-sticks error rows once an admin "
              "re-assigns the dataset.", {})])

h2("6.5  Metadata sync — full and delta")
para("MetadataSyncService is the only place in the app that holds an API client for "
     "metadata; every MetadataResource is constructed with just a database handle, "
     "so by construction it cannot reach the network on its own.")
bullet_rich([("Full sync", {'bold': True}),
             (" (first login into an empty database): clears all metadata / link "
              "tables, then syncs every resource in dependency order — category "
              "options → categories → combos → option combos → option sets → "
              "options → attributes → data elements → indicators → data-element "
              "groups → organisation units → data sets → sections → validation "
              "rules. The lastMetadataSync timestamp is written only after the full "
              "pass succeeds, so an interrupted download retries a full sync rather "
              "than getting stuck on the delta path with half-empty metadata.", {})])
bullet_rich([("Delta sync", {'bold': True}),
             (" (every subsequent online login): each resource fetches a cheap "
              "id + lastUpdated list, compares against local pairs, fetches only "
              "changed/new objects (chunked at 100 IDs per request), and deletes "
              "locally anything the server no longer has.", {})])
bullet_rich([("Org-unit scoping. ", {'bold': True}),
             ("A facility user's metadata sync pulls only their own capture-root "
              "subtree, never the ~38,000-unit national tree. For offline storage "
              "the subtree is further bounded to the capture roots plus their direct "
              "children — a woreda-assigned user stores the woreda and its PHCUs, "
              "not every health post beneath them. Online, the full subtree stays "
              "reachable via a live fallback. Delta sync converges a device that "
              "synced before this bound existed down to the bound.", {})])
bullet_rich([("Failure model. ", {'bold': True}),
             ("Each resource syncs in its own step; a drop mid-sequence leaves "
              "earlier resources fully synced and later ones untouched — stale but "
              "internally consistent, never half-written.", {})])

h2("6.6  Background-kill protection (Android)")
para("Once the app is backgrounded, Android's execution limits and — more "
     "aggressively — several OEM battery managers will kill the process mid-push. "
     "Two mechanisms mitigate this:")
bullet_rich([("SyncForegroundService", {'bold': True}),
             (" — a native Android foreground service (dataSync type) started and "
              "stopped by the SyncCoordinator around each push attempt, so the "
              "notification only appears when there is actually work to protect. "
              "No-op on iOS (which caps background execution at a few seconds "
              "regardless).", {})])
bullet_rich([("BatteryOptimization", {'bold': True}),
             (" — a one-time (ever) prompt after login asking the user to exempt "
              "the app from OEM battery optimisation. The exemption still requires "
              "explicit user consent via the system dialog.", {})])

h2("6.7  Chart draft coordinator")
para("A local dashboard saved while offline is stored as a draft; "
     "ChartDraftCoordinator (started in main) finishes those drafts — running their "
     "analytics query and caching the result — the moment connectivity returns.")

h2("6.8  Schema for a locally-queued write table")
para("The invariant every sync-aware class relies on: a new write-queue table "
     "carries syncState (the four-value enum), syncError (nullable text), and "
     "lastModified (sourced from PeriodAccess.effectiveNow(), never raw "
     "DateTime.now()). A transport failure leaves a row exactly where it was; only a "
     "server verdict (2xx accepted or 409 rejected) may change a row's sync state.")
doc.add_page_break()

# ===========================================================================
# 7. VISUALIZATION
# ===========================================================================
h1("7  Visualization and Analytics")

h2("7.1  Three areas")
para("The Visualization side of the Home toggle presents three flat tabs:")
table(["Tab", "What it is", "Connectivity"],
      [["Server Dashboard", "DHIS2 server dashboards, rendered natively with "
        "fl_chart via the same /api/dashboards → /api/visualizations/{id} → "
        "/api/analytics pipeline the DHIS2 web app uses. Each visualization on a "
        "dashboard is fetched and drawn independently, so one slow or broken chart "
        "never blocks the rest.",
        "Online to load; last successful result cached for offline viewing."],
       ["Local Dashboard", "Charts the user has built and saved on this device. The "
        "saved configuration lives as JSON under one key in the per-user database "
        "(no schema change); every successful query result is cached under its own "
        "key so the chart is still viewable offline. Full CRUD — create, view, "
        "edit, delete — with no server push.",
        "Online to (re)run a chart; cached result offline."],
       ["Create New", "The chart builder: pick indicators or data elements (by "
        "group, or flat from local metadata when offline), an org unit, a period, "
        "and a chart type. Saving lands the chart in Local Dashboard.",
        "Dimension pickers prefer the live server, fall back to local metadata."]],
      widths=[1.2, 3.6, 1.4])

h2("7.2  Defensive behaviour")
bullet("A visualization with no period dimension anywhere (columns / rows / "
       "filters) is refused before the analytics call — DHIS2 would answer such a "
       "query with a permanent 409 — and a specific, displayable "
       "“misconfigured visualization” message is shown instead of a "
       "generic network error.")
bullet("Relative-period flags (last12Months, …) are translated to their analytics "
       "dimension IDs (LAST_12_MONTHS) mechanically.")
bullet("A dashboard tracks how many item types it cannot render yet (MAP, TEXT, "
       "EVENT_CHART, …) so the UI can say “2 items not supported” rather "
       "than silently dropping them.")
bullet("An offline-cache banner tells the user when a chart is showing cached "
       "rather than live data.")
doc.add_page_break()

# ===========================================================================
# 8. REMINDERS & ONBOARDING
# ===========================================================================
h1("8  Reminders, To-Do, and Onboarding")

h2("8.1  Expected reports (the “To-do” band)")
rich([("CaptureRepository.getExpectedReports() ", {'mono': True}),
      ("computes, on the fly, every report the user still owes: each dataset "
       "assigned to one of their facilities (capture roots + direct children), for "
       "every period still open for entry, that is not already completed locally or "
       "on the server. It best-effort pulls the server's completion state first when "
       "online, so a report finished on the web does not linger. Results are sorted "
       "most-urgent first.", {})])
table(["Urgency", "Meaning"],
      [["overdue", "The period has ended but is still inside its lock window — "
        "fill it now."],
       ["dueSoon", "The period is ending within a few days."],
       ["open", "Open, deadline not yet close."]],
      widths=[1.1, 5.1])

h2("8.2  Deadline reminders")
para("ReportReminderService schedules on-device local notifications (no server, no "
     "push) for reports due before their lock date. Rebuilt from scratch every time "
     "the expected-reports list is recomputed, so they always match reality.")
bullet("Grouped by lock date — a facility owing five reports on the same day gets "
       "one notification, not five.")
bullet("At most two per day: a heads-up three days out and a “locks "
       "today” on the day, both at 08:00.")
bullet_rich([("Time zone is pinned to Africa/Addis_Ababa", {'bold': True}),
             (" — this is a Federal Ministry of Health app and every period boundary "
              "runs on Addis time.", {})])
bullet("Inexact alarms only (AndroidScheduleMode.inexactAllowWhileIdle) — a "
       "reminder a few minutes late is fine, and this avoids the "
       "SCHEDULE_EXACT_ALARM permission entirely.")
bullet("No-op on web; self-disables silently if the platform refuses "
       "initialisation or permission.")

h2("8.3  Onboarding and app tour")
bullet("A first-run carousel (gated by a SharedPreferences flag, independent of "
       "login/logout).")
bullet_rich([("Per-screen spotlight tours", {'bold': True}),
             (" (showcaseview) for Home, routine data entry, disease data entry, "
              "and Visualization — each with its own ID so it triggers once, on "
              "that screen's first visit, rather than one monolithic tour spanning "
              "navigations.", {})])
bullet("“Take the tour again” in the Home drawer clears every screen's "
       "tour flag; Home's replays immediately, the others on next visit.")
doc.add_page_break()

# ===========================================================================
# 9. SECURITY POSTURE
# ===========================================================================
h1("9  Security Posture")

h2("9.1  Authentication")
para("DHIS2 uses HTTP Basic authentication (Authorization: Basic "
     "base64(username:password)). The app builds that header in exactly one place. "
     "SessionService.login() is the single decision tree for both login paths and "
     "returns one of: onlineFirstSync, onlineReturning, offline, offlineNoCache, "
     "invalidCredentials.")
bullet_rich([("Online path: ", {'bold': True}),
             ("a throwaway credentialed Dio client calls GET /api/me.json; 200 → "
              "valid, 401 → invalid. On success: open the per-user database; anchor "
              "the tamper-resistant clock to the server Date header (only if no "
              "local data is pending); persist the offline verifier; kick off "
              "metadata sync (full on first login, delta after) in the background.", {})])
bullet_rich([("Offline path: ", {'bold': True}),
             ("the per-user database must already exist (otherwise offlineNoCache — "
              "there is no way to log in for the very first time without a "
              "connection). The verifier is recomputed and compared in constant "
              "time; a local backward-clock check runs.", {})])

h2("9.2  The offline verifier")
code_block("verifier = SHA-256( salt : serverUrl : username : password )")
bullet("salt is 16 cryptographically random bytes (Random.secure()), unique per "
       "user, stored alongside the hash.")
bullet_rich([("serverUrl is normalised and baked into the hash — ", {}),
             ("a verifier created against one DHIS2 instance cannot validate a login "
              "aimed at another.", {'bold': True})])
bullet("Comparison is constant-time (XOR-accumulate over every byte) to resist "
       "timing attacks.")
bullet_rich([("The password itself is never stored, in any form, online or "
              "offline.", {'bold': True}),
             (" The verifier can check a login attempt; it cannot be reversed into "
              "a password.", {})])
bullet("Both salt and hash are held in flutter_secure_storage, so the verifier is "
       "encrypted at rest even though it is already a one-way hash.")

h2("9.3  Credential and session storage")
table(["Store", "Holds", "Backing"],
      [["CredentialStore", "The offline-login verifier only (salt + hash, per user).",
        "flutter_secure_storage"],
       ["SecureStorage", "Session / profile convenience data: username, base URL, "
        "cached /me JSON, org units. Plus a legacy auth_token key that is only ever "
        "deleted (on 401 / logout), never written.",
        "flutter_secure_storage"],
       ["SQLite database", "Cached metadata, pending field data, local audit log, "
        "local dashboard configs. The UsersTable caches only uid / username / "
        "displayName from /api/me.",
        "Plain SQLite file (not encrypted at rest)"]],
      widths=[1.3, 3.5, 1.4])
rich([("Platform backing for secure storage: ", {}),
      ("Android → EncryptedSharedPreferences (Android Keystore-backed); "
       "iOS → Keychain (first-unlock accessibility). ", {}),
      ("Nothing sensitive lives in the SQLite database — no password, no token, no "
       "verifier.", {'bold': True})])

h2("9.4  Logout, wipe, and 401 handling")
bullet_rich([("Logout", {'bold': True}),
             (" closes the database connection and clears the in-memory session but "
              "keeps the database file and the offline verifier — the same user can "
              "log back in offline later without re-downloading metadata. The "
              "logout confirmation dialog explicitly tells the user that unsynced "
              "data stays on the device.", {})])
bullet_rich([("Wipe", {'bold': True}),
             (" deletes the database file (+ WAL/SHM sidecars) and the verifier. "
              "Guarded: it throws unless confirmedDataLoss is passed, after checking "
              "the count of unsynced work — callers must show that count and get "
              "explicit confirmation.", {})])
bullet_rich([("401 anywhere", {'bold': True}),
             (" (AuthInterceptor): deletes the legacy token, ends the session "
              "(not a wipe), and redirects to /login?reason=session-expired. The "
              "local database and verifier survive — the user can still log in "
              "offline.", {})])

h2("9.5  Transport and platform")
bullet("HTTPS-only in release. There is no usesCleartextTraffic flag in the main "
       "Android manifest; the debug manifest allows cleartext explicitly, for local "
       "dev servers. iOS App Transport Security applies unmodified (no exceptions in "
       "Info.plist).")
bullet("Dio timeouts: 30 s each for connect / receive / send.")
bullet_rich([("Release signing is enforced, not optional. ", {'bold': True}),
             ("The Android build hard-fails the moment a *Release task is scheduled "
              "if android/key.properties is missing — this prevents an "
              "accidentally debug-signed release that can never be updated in place "
              "(a user would have to uninstall, losing unsynced data). An explicit "
              "opt-in flag exists for non-distributable local test builds.", {})])
bullet("SQL injection is eliminated by construction — all queries go through Drift's "
       "type-safe builder; the only raw statements are fixed PRAGMAs with no user "
       "input.")

h2("9.6  Android permissions")
para("Every permission the app requests, and why:")
table(["Permission", "Purpose"],
      [["INTERNET", "API calls."],
       ["ACCESS_NETWORK_STATE", "Connectivity detection."],
       ["FOREGROUND_SERVICE, FOREGROUND_SERVICE_DATA_SYNC", "Keep an in-flight "
        "offline-data push alive after the app is backgrounded."],
       ["POST_NOTIFICATIONS", "Report deadline reminders (and the sync foreground "
        "notification)."],
       ["RECEIVE_BOOT_COMPLETED", "Re-arm deadline reminders after a reboot. "
        "Inexact alarms only — no SCHEDULE_EXACT_ALARM."],
       ["REQUEST_IGNORE_BATTERY_OPTIMIZATIONS", "Let the user exempt the app from "
        "OEM battery-optimisation killing. Install-time permission; the exemption "
        "itself still needs explicit user consent via the system dialog."]],
      widths=[2.7, 3.5])
para("No storage, camera, location, contacts, or exact-alarm permissions are "
     "requested anywhere.", italic=True)

h2("9.7  Logging discipline")
bullet("Release builds log warning and error only — request/response detail never "
       "reaches a field device's logcat.")
bullet("The Authorization header is explicitly redacted before any logging, and the "
       "credentialed API-client variant carries a hard rule never to log headers.")
bullet("Response bodies are logged only at debug level, truncated — and debug "
       "logging is stripped from release entirely.")

h2("9.8  Honest summary")
table(["Area", "Status"],
      [["Password storage", "OK — never stored in any recoverable form, online or "
        "offline"],
       ["Offline verification", "OK — salted SHA-256, constant-time compare, "
        "server-URL-bound"],
       ["Secure storage", "OK — OS-level (Android Keystore / iOS Keychain)"],
       ["Transport", "OK — HTTPS-only in release  //  gap: no certificate pinning"],
       ["Local database at rest", "Gap — not encrypted (holds cached metadata + "
        "pending field data + local audit log; never credentials)"],
       ["Android release build", "OK — hard-fails without a real keystore  //  "
        "gaps: no R8/ProGuard minification; android:allowBackup not explicitly "
        "disabled"],
       ["Logging", "OK — Authorization redacted; release suppresses debug/info"],
       ["SQL injection", "OK — eliminated by construction (Drift only)"],
       ["Session on 401", "OK — ends session, redirects with a reason, preserves "
        "local data"]],
      widths=[1.9, 4.3])
para("None of the gaps expose credentials or allow remote data tampering. They are "
     "defense-in-depth items for the “device is lost or physically "
     "compromised” threat model, each with a concrete, scoped fix (section 13).")
doc.add_page_break()

# ===========================================================================
# 10. DATABASE SCHEMA
# ===========================================================================
h1("10  Database Schema")

h2("10.1  Technology")
para("Drift over SQLite. Type-safe, compile-time-checked queries; row and companion "
     "classes are code-generated from Table subclasses — there is no hand-written "
     "SQL-to-Dart mapping anywhere. Native builds use sqlite3_flutter_libs; web uses "
     "sqlite3.wasm via a Drift worker. Journal mode is WAL; foreign keys are on.")

h2("10.2  Tables (30, current schema version 5)")
h3("Metadata tables (13) — mirror DHIS2 configuration, cleared and refreshed by a full sync")
para("OrgUnitsTable, DataSetsTable, DataElementsTable, SectionsTable, "
     "IndicatorsTable, CategoriesTable, CategoryOptionsTable, CategoryCombosTable, "
     "CategoryOptionCombosTable, OptionSetsTable, OptionsTable, "
     "DataElementGroupsTable, ValidationRulesTable.")
h3("Link / join tables (10) — n:n relationships")
para("DataSetElementsTable (carries sortOrder, compulsory, and an effective "
     "categoryCombo override), CompulsoryDataElementOperandsTable, "
     "DataSetOrgUnitsTable, SectionDataElementsTable, SectionIndicatorsTable, "
     "SectionGreyFieldsTable, DataElementGroupMembersTable, "
     "CategoryCategoryOptionsTable, CategoryComboCategoriesTable, "
     "CategoryOptionComboOptionsTable.")
h3("Global / control tables (4)")
bullet("UsersTable — the cached /api/me response (enables offline getCurrentUser and "
       "offline-login user reconstruction).")
bullet("AttributesTable / AttributeValuesTable — a generic custom-attribute "
       "mechanism keyed on (objectType, objectUid, attributeUid), so any metadata "
       "type can carry attribute values without a table-per-type explosion.")
bullet("SyncInfoTable — a key-value store for sync bookkeeping (lastMetadataSync, "
       "clock high-water mark, tamper timestamp), the outlier history cache, and "
       "local dashboard configs + result caches.")
h3("Data / write-queue tables (3) — the only tables a metadata sync never clears")
bullet("DataValuesTable — composite key (dataElementUid, period, orgUnitUid, "
       "categoryOptionComboUid, attributeOptionComboUid); value, comment, storedBy, "
       "syncState, syncError, lastModified.")
bullet("CompleteDataSetRegistrationsTable — composite key (dataSetUid, period, "
       "orgUnitUid, attributeOptionComboUid); completed, syncState, syncError, "
       "date, lastModified.")
bullet("AuditLogTable — append-only local edit history (entityType, auditType, "
       "the entity keys, previousValue, newValue, modifiedBy, modifiedAt).")

h2("10.3  Migrations")
rich([("The schema is no longer frozen. ", {'bold': True}),
      ("schemaVersion is 5. onUpgrade runs one ordered step per version bump; a "
       "missing step throws loudly rather than opening a field device's database "
       "against a schema it does not match (which would corrupt real data).", {})])
table(["Version", "Change"],
      [["1", "Frozen baseline."],
       ["2", "Create AuditLogTable (local edit history)."],
       ["3", "Add AuditLogTable.auditType (explicit CREATE/UPDATE/DELETE); backfill "
        "existing rows from previousValue/newValue nullness."],
       ["4", "Add DataSetElementsTable.compulsory — mandatory-field support."],
       ["5", "Create CompulsoryDataElementOperandsTable — the element+combo-pair "
        "sibling of dataSetElement.compulsory."]],
      widths=[0.8, 5.4])
para("Steps 3 and 4 guard against a device that crashed mid-migration by checking "
     "column existence first, so a retry does not fail forever on “duplicate "
     "column”. Adding a table means: define it, add it to the "
     "@DriftDatabase list, add it to the metadata-clear list if it is metadata, "
     "bump schemaVersion with a step, and regenerate.")

h2("10.4  Testing against the database")
para("Tests use an in-memory database (AppDatabase.forTesting(NativeDatabase."
     "memory())) — no file I/O, no platform plugins — which is what lets the full "
     "test suite run in plain CI with no emulator.")
doc.add_page_break()

# ===========================================================================
# 11. BUILD, RELEASE & CI
# ===========================================================================
h1("11  Build, Release, and CI")

h2("11.1  Continuous integration")
para(".github/workflows/ci.yml runs on every push to main / integration* and every "
     "pull request: flutter pub get → flutter analyze → flutter test. Pinned to "
     "Flutter 3.41.4. A red CI blocks merge. There is no build/deploy job in the "
     "workflow — analyze + test only.")

h2("11.2  Android release build")
para("Configured in android/app/build.gradle.kts. Release builds require a real "
     "signing keystore and the build hard-fails without one.")
numbered("Generate a keystore once (keytool -genkey …). Guard the file and password "
         "— an Android app can only ever be updated by APKs signed with the same "
         "key.")
numbered("Create android/key.properties (gitignored) with storeFile / storePassword "
         "/ keyPassword / keyAlias.")
numbered("flutter build apk --release.")
para("A non-distributable debug-signed test APK can be built with an explicit opt-in "
     "(ALLOW_DEBUG_SIGNING=true), which logs a loud warning. Not yet configured: "
     "R8/ProGuard minification; android:allowBackup=\"false\".")

h2("11.3  Web build")
para("flutter build web. The WASM database assets (sqlite3.wasm, drift_worker.js) "
     "are served from web/ and are version-pinned to the drift / sqlite3 package "
     "versions — re-download them together when bumping either. For local "
     "development against staging use ./run_web.sh, which pins the dev server to "
     "port 3001 (the only port besides 3000 that staging's CORS policy allows).")

h2("11.4  Code generation")
para("The Drift database uses generated code (app_database.g.dart, ~16,900 lines, "
     "committed). Regenerate only after editing a table definition: "
     "dart run build_runner build --delete-conflicting-outputs.")
doc.add_page_break()

# ===========================================================================
# 12. TESTING
# ===========================================================================
h1("12  Testing")
para("24 test files, ~5,200 lines, mirroring lib/ under test/. Coverage is "
     "deliberately weighted toward the offline-critical core — the code where a bug "
     "would silently corrupt or lose field data, and the code hardest to verify by "
     "clicking through the UI.")
h3("Well covered")
bullet("Sync engine and conflict resolution (data_value_store, data_value_sync, "
       "data_value_push, completeness + completeness pull).")
bullet("Data quality — value-type validator, validation service (rule evaluation), "
       "outlier detection, controller-element service, element/indicator label "
       "services.")
bullet("Ethiopian calendar and period service.")
bullet("Metadata resources (data set, organisation unit, section) and Drift "
       "round-tripping.")
bullet("Capture repository (including expected-reports computation), local "
       "visualization repository, disease-entry list, PDF export, date-filter "
       "resolution.")
h3("Thin")
bullet("Bloc-level tests (AuthBloc, DataEntryBloc have no dedicated files) and full "
       "widget / integration tests beyond smoke-level rendering.")
h3("Conventions")
bullet("In-memory database for anything touching AppDatabase.")
bullet("Network-touching classes take an ApiClient by constructor injection; tests "
       "supply a canned/fake HTTP adapter. No mockito / mocktail — the project "
       "favours real fakes and in-memory implementations.")
doc.add_page_break()

# ===========================================================================
# 13. ROADMAP & KNOWN LIMITATIONS
# ===========================================================================
h1("13  Roadmap and Known Limitations")
para("Stated transparently. None of these involve credential exposure or remote "
     "data tampering.")

h2("13.1  Security hardening (before large-scale rollout)")
numbered("Encrypt the local database at rest (SQLCipher-backed SQLite). It holds "
         "cached metadata, pending field data, and the local audit log — never "
         "credentials — but a lost device should still not expose facility health "
         "data in a plain-readable file.")
numbered("Enable R8/ProGuard minification for Android release builds.")
numbered("Certificate pinning for the production DHIS2 host.")
numbered("Set android:allowBackup=\"false\" — health data should not ride Android's "
         "default auto-backup.")

h2("13.2  Engineering investment")
numbered("Bloc and widget test coverage (the offline-critical core is well tested; "
         "the Blocs and most widgets are not).")
numbered("A lightweight service locator (get_it) to remove the repeated manual-DI "
         "wiring across pages.")
numbered("Crash and performance monitoring (there is currently no automated crash "
         "reporting).")

h2("13.3  Product roadmap (depends on Ministry / HISP priorities)")
bullet("Biometric app-lock for at-rest device access.")
bullet("Role-based dashboards; a fuller in-app audit view.")
bullet("A move from HTTP Basic to OAuth2 / PAT token auth once the target DHIS2 "
       "deployment supports it.")
bullet("Multi-language support (Amharic, Afaan Oromo, Tigrinya) — the app is "
       "currently English-only.")
bullet("Dark mode.")

h2("13.4  Known, accepted trade-offs")
bullet("Manual dependency injection instead of a framework — simple and traceable, "
       "at the cost of repeated wiring code.")
bullet("The data-entry UI does not support non-default attribute category combos "
       "for Routine datasets — every value is stored under the dataset instance's "
       "default combo (Disease Registration datasets do use a real combo, picked "
       "before the form opens). A scope decision, not a bug.")
bullet("Server-side validation rule features beyond basic arithmetic (d2 functions, "
       "constants, indicators, org-unit groups) are skipped by the offline "
       "evaluator rather than approximated.")
bullet("Indicator groups are not synced for offline use (only indicators "
       "themselves), so the offline chart builder falls back to a flat indicator "
       "list.")
bullet("OFFLINE_INTEGRATION.md at the repo root is a pre-implementation design "
       "note, superseded by the current implementation — kept for historical "
       "context only.")

h2("13.5  Staging-environment issues (not app bugs)")
bullet_rich([("The staging server's own “current period” has been observed "
              "running ~2 months behind the real Ethiopian calendar", {'bold': True}),
             (" (investigated 2026-07-21). Data pushed “for today” lands "
              "correctly but the staging web UI will not display that period for "
              "weeks. The app's calendar conversion is the correct one — verified "
              "two independent ways; the server's period generation is what is "
              "behind. Check via the API directly before assuming a sync bug.", {})])
bullet("The staging “Routine Data Entry” web app's section tabs do not "
       "switch past the first — a front-end bug in that app, unrelated to this "
       "codebase.")
doc.add_page_break()

# ===========================================================================
# APPENDIX A — GLOSSARY
# ===========================================================================
h1("Appendix A  Glossary")
table(["Term", "Meaning"],
      [["DHIS2", "District Health Information System 2 — the platform this app feeds."],
       ["Aggregate data", "Counts and totals for a period (as opposed to "
        "individual-level tracker data). This app is aggregate-only."],
       ["Data element", "One thing being counted (e.g. “ANC 1st visit”)."],
       ["Category option combo (COC)", "A disaggregation of a data element (e.g. "
        "“<1 year, Male”). Forms have data elements as rows and COCs as "
        "columns."],
       ["Attribute option combo (AOC)", "A whole-form disaggregation (e.g. a funding "
        "stream, or Disease Registration's Department × Outcome)."],
       ["Data set", "A form: a set of data elements captured together for a period "
        "and org unit."],
       ["Organisation unit (org unit)", "A node in the health facility hierarchy "
        "(country → region → zone → woreda → facility)."],
       ["Period", "The reporting window (a month, quarter, year, …), identified by a "
        "DHIS2 period ID."],
       ["expiryDays", "Days after a period ends during which the server still "
        "accepts data for it. 0 = never locks."],
       ["openFuturePeriods", "How many not-yet-started periods the server will "
        "accept data for."],
       ["Completion registration", "A record that a form was signed off as complete "
        "for a period / org unit."],
       ["Validation rule", "A server-defined cross-check between data elements "
        "(e.g. A ≤ B). Warns, does not block."],
       ["Grey field", "A cell the server configures as disabled for a section."],
       ["Drift", "The Dart persistence library wrapping SQLite used here."],
       ["Bloc", "The state-management pattern (events in, states out) used "
        "throughout."]],
      widths=[1.9, 4.3])

# ===========================================================================
# APPENDIX B — KEY FILE REFERENCE
# ===========================================================================
h1("Appendix B  Key File Reference")
para("A newcomer's map to where each concern actually lives.")
table(["Concern", "File(s)"],
      [["App entry / startup", "lib/main.dart"],
       ["Login decision tree", "lib/core/auth/session_service.dart"],
       ["Offline verifier", "lib/core/auth/credential_store.dart"],
       ["Session + router guard", "lib/core/auth/app_session.dart, "
        "lib/core/router/app_router.dart"],
       ["HTTP client + interceptors", "lib/core/network/api_client.dart, "
        "lib/core/network/interceptors/"],
       ["Connectivity", "lib/core/network/network_info.dart, "
        "lib/core/network/connectivity_service.dart"],
       ["Database schema + migrations", "lib/core/database/app_database.dart"],
       ["Metadata sync", "lib/core/metadata/metadata_sync_service.dart, "
        "lib/core/metadata/metadata_resource.dart"],
       ["Sync engine", "lib/core/sync/ (drift_sync_manager, sync_coordinator, "
        "manual_sync, sync_foreground_service, battery_optimization)"],
       ["Write path + conflict resolution", "lib/core/data/data_value_store.dart, "
        "data_value_sync.dart, data_value_push.dart, completeness.dart"],
       ["Tamper-resistant clock", "lib/core/data/period_access.dart"],
       ["Ethiopian calendar", "lib/core/data/ethiopian_calendar.dart, "
        "ethiopian_period_service.dart"],
       ["Value-type validation", "lib/core/data/value_type_validator.dart"],
       ["Validation rules (offline)", "lib/core/data/validation_service.dart"],
       ["Outlier detection", "lib/core/data/outlier_detection_service.dart, "
        "lib/features/data_entry/domain/entities/outlier_stats.dart"],
       ["Controller elements", "lib/core/data/controller_element_service.dart"],
       ["Audit trail", "lib/core/data/audit_log_store.dart, server_audit_service.dart, "
        "lib/features/audit_log/"],
       ["Deadline reminders", "lib/core/notifications/report_reminder_service.dart"],
       ["Onboarding / tour", "lib/core/onboarding/, lib/features/onboarding/"],
       ["Capture flow", "lib/features/capture/"],
       ["Entry form", "lib/features/data_entry/"],
       ["Visualization", "lib/features/visualization/"],
       ["Developer docs", "docs/ (16+ chapters, file-cited)"]],
      widths=[1.9, 4.3])

hr()
para("End of document. Prepared from the repository at the current development head, "
     f"{LONG_DATE}. For the most current behaviour of any specific claim, "
     "verify against the file cited.", italic=True, size=9, color=GREY)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out_dir = os.path.join(REPO, 'technical_documentation')
os.makedirs(out_dir, exist_ok=True)
out = os.path.join(out_dir, 'RDHIS2_Mobile_Technical_Documentation.docx')
doc.save(out)
print("saved", out)

soffice = shutil.which('soffice') or shutil.which('libreoffice')
if soffice:
    subprocess.run([soffice, '--headless', '--convert-to', 'pdf',
                    '--outdir', out_dir, out], check=False)
    print("saved", out.replace('.docx', '.pdf'))
else:
    print("libreoffice not found — skipping PDF conversion")
